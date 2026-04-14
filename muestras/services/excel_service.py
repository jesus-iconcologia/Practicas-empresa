import io
import openpyxl
from openpyxl.styles import PatternFill


def cargar_workbook_desde_path(path):
    return openpyxl.load_workbook(path)


def workbook_to_response_bytes(workbook):
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def aplicar_color_celda(celda, color_hex):
    celda.fill = PatternFill("solid", fgColor=color_hex)