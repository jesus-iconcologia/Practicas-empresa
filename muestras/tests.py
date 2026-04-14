import io
import json
import zipfile
import os
import shutil
import tempfile

import openpyxl
from datetime import date
from urllib.parse import quote
from django.contrib.auth.models import Group, Permission, User
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase
from django.test.utils import override_settings

from .forms import CongeladorForm, ReparacionRecursoForm
from .models import (
    Bandeja,
    Caja,
    Congelador,
    Documento,
    Estante,
    Estudio,
    Muestra,
    Rack,
    ReparacionRecurso,
    Subposicion,
    agenda_envio,
    registro_destruido,
)


TEST_MEDIA_ROOT = tempfile.mkdtemp(prefix="geslab-test-media-")


@override_settings(MEDIA_ROOT=TEST_MEDIA_ROOT)
class GesLabRegressionTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(TEST_MEDIA_ROOT, ignore_errors=True)

    def _parse_streaming_json_response(self, response):
        payload = b"".join(response.streaming_content).decode("utf-8").strip().splitlines()
        self.assertTrue(payload)
        return json.loads(payload[-1])

    def setUp(self):
        self.user = User.objects.create_user(username="tester", password="secret")
        self.user.is_superuser = True
        self.user.is_staff = True
        self.user.save()
        self.client.force_login(self.user)

        self.congelador = Congelador.objects.create(
            congelador="FZ-01",
            tipo_estructura=Congelador.ESTRUCTURA_CONGELADOR_80,
        )
        self.estante = Estante.objects.create(congelador=self.congelador, numero="1")
        self.rack = Rack.objects.create(estante=self.estante, numero="A", posicion_rack_estante="1")
        self.bandeja = Bandeja.objects.create(rack=self.rack, numero="B1", posicion_bandeja_rack="1")
        self.caja = Caja.objects.create(rack=self.rack, bandeja=self.bandeja, numero="C1", posicion_caja_rack="1")
        self.subposicion_libre = Subposicion.objects.create(
            caja=self.caja,
            numero="A-1",
            fila="A",
            columna="1",
            vacia=True,
        )
        self.congelador_20 = Congelador.objects.create(
            congelador="FZ-20",
            tipo_estructura=Congelador.ESTRUCTURA_CONGELADOR_20,
        )
        self.estante_20 = Estante.objects.create(congelador=self.congelador_20, numero="E1")
        self.cajon_20 = Rack.objects.create(
            estante=self.estante_20,
            numero="C1",
            posicion_rack_estante="1",
        )
        self.caja_20 = Caja.objects.create(
            rack=self.cajon_20,
            bandeja=None,
            numero="BOX-20",
            posicion_caja_rack="",
        )
        self.subposicion_20 = Subposicion.objects.create(
            caja=self.caja_20,
            numero="A-1",
            fila="A",
            columna="1",
            vacia=True,
        )
        self.nevera = Congelador.objects.create(
            congelador="NV-01",
            tipo_estructura=Congelador.ESTRUCTURA_NEVERA,
        )
        self.estante_nevera = Estante.objects.create(congelador=self.nevera, numero="N1")
        self.cajon_nevera = Rack.objects.create(
            estante=self.estante_nevera,
            numero="CJ1",
            posicion_rack_estante="2",
        )
        self.caja_nevera = Caja.objects.create(
            rack=self.cajon_nevera,
            bandeja=None,
            numero="BOX-NV",
            posicion_caja_rack="",
        )
        self.subposicion_nevera = Subposicion.objects.create(
            caja=self.caja_nevera,
            numero="B-2",
            fila="B",
            columna="2",
            vacia=True,
        )

    def test_api_subposiciones_devuelve_estado_legible(self):
        muestra = Muestra.objects.create(
            nom_lab="LAB-001",
            estado_actual="PENV",
        )
        self.subposicion_libre.muestra = muestra
        self.subposicion_libre.vacia = False
        self.subposicion_libre.save()

        response = self.client.get(f"/api/get_subposiciones_por_caja_tree/?caja_id={self.caja.id}")

        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["subposiciones"][0]["muestra_estado"], "Parcialmente enviada")
        self.assertEqual(data["subposiciones"][0]["muestra_detalle_url"], "/archivo/detalles_muestra/LAB-001")

    def test_historial_localizaciones_muestra_destruida_muestra_fila_final(self):
        muestra = Muestra.objects.create(
            nom_lab="LAB-DEST",
            estado_actual="DEST",
        )
        registro_destruido.objects.create(muestra=muestra, usuario=self.user)

        response = self.client.get(f"/muestras/historial_localizaciones/{muestra.id}")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Destruida")

    def test_localizaciones_muestra_resumen_por_recurso(self):
        muestra = Muestra.objects.create(
            nom_lab="LAB-SUM",
            estado_actual="DISP",
        )
        self.subposicion_libre.muestra = muestra
        self.subposicion_libre.vacia = False
        self.subposicion_libre.save()

        response = self.client.get("/archivo/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Recurso FZ-01")
        self.assertNotContains(response, "1 posiciones")
        self.assertNotContains(response, "0 libres")
        self.assertNotContains(response, "1 ocupadas")
        self.assertNotContains(response, "1 cajas")

    def test_localizaciones_con_filtro_caja_mantiene_la_jerarquia_renderizada(self):
        response = self.client.get(f"/archivo/?filtro_caja={self.caja.id}")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, f'data-caja-id="{self.caja.id}"')
        self.assertContains(response, "Caja C1")

    def test_descarga_excel_envios_excluye_muestras_destruidas(self):
        centro = agenda_envio.objects.create(
            centro="Centro X",
            lugar="Madrid",
            direccion="Calle Mayor 1",
        )
        muestra_valida = Muestra.objects.create(
            nom_lab="LAB-OK",
            estado_actual="DISP",
            volumen_actual=10,
            unidad_volumen="mL",
            concentracion_actual=5,
            unidad_concentracion="mg/mL",
        )
        muestra_destruida = Muestra.objects.create(
            nom_lab="LAB-DEST",
            estado_actual="DEST",
            volumen_actual=1,
            unidad_volumen="mL",
            concentracion_actual=1,
            unidad_concentracion="mg/mL",
        )
        session = self.client.session
        session["muestras_envio"] = [muestra_valida.id, muestra_destruida.id]
        session.save()

        response = self.client.post(
            f"/muestras/envio/agenda/{centro.id}/envio/upload_excel_envio",
            {"descargar_excel_envio": "1"},
        )

        self.assertEqual(response.status_code, 200)
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        sheet = workbook.active
        muestras_exportadas = [sheet.cell(row=row, column=1).value for row in range(2, sheet.max_row + 1)]

        self.assertIn("LAB-OK", muestras_exportadas)
        self.assertNotIn("LAB-DEST", muestras_exportadas)

    def test_confirmar_importacion_estudios_crea_registro(self):
        session = self.client.session
        session["filas_validas"] = [
            {
                "referencia_estudio": "REF-001",
                "nombre_estudio": "ESTUDIO-IMPORT",
                "descripcion_estudio": "Importado desde test",
                "fecha_inicio_estudio": "2026-01-10",
                "fecha_fin_estudio": "2026-02-10",
                "investigador_principal": "Dra. Test",
            }
        ]
        session.save()

        response = self.client.post("/estudios/excel", {"confirmar": "1"})

        self.assertEqual(response.status_code, 200)
        data = self._parse_streaming_json_response(response)
        self.assertEqual(data["status"], "done")
        self.assertTrue(Estudio.objects.filter(nombre_estudio="ESTUDIO-IMPORT").exists())

    def test_bootstrap_initial_data_crea_grupos_base(self):
        from django.contrib.auth.models import Group

        Group.objects.filter(name__in=["Tecnicos de laboratorio", "Investigadores"]).delete()

        call_command("bootstrap_initial_data")

        tecnicos = Group.objects.get(name="Tecnicos de laboratorio")
        investigadores = Group.objects.get(name="Investigadores")

        self.assertTrue(tecnicos.permissions.filter(codename="can_view_muestras_web").exists())
        self.assertTrue(tecnicos.permissions.filter(codename="can_view_localizaciones_web").exists())
        self.assertTrue(tecnicos.permissions.filter(codename="can_view_estudios_web").exists())
        self.assertTrue(investigadores.permissions.filter(codename="can_view_estudios_web").exists())
        self.assertTrue(investigadores.permissions.filter(codename="can_change_estudios_web").exists())

    def test_eliminar_documento_borra_todos_los_seleccionados(self):
        estudio = Estudio.objects.create(nombre_estudio="EST-001")
        doc1 = Documento.objects.create(
            estudio=estudio,
            archivo=SimpleUploadedFile("doc1.txt", b"uno"),
            usuario_subida=self.user,
        )
        doc2 = Documento.objects.create(
            estudio=estudio,
            archivo=SimpleUploadedFile("doc2.txt", b"dos"),
            usuario_subida=self.user,
        )
        session = self.client.session
        session["id"] = estudio.id
        session.save()

        response = self.client.post(
            "/estudios/eliminar_documento",
            {"doc_id": [str(doc1.id), str(doc2.id)]},
        )

        self.assertEqual(response.status_code, 302)
        self.assertFalse(Documento.objects.filter(id=doc1.id).exists())
        self.assertFalse(Documento.objects.filter(id=doc2.id).exists())

    def test_descargar_documentos_seleccionados_devuelve_zip(self):
        estudio = Estudio.objects.create(nombre_estudio="EST-ZIP")
        doc1 = Documento.objects.create(
            estudio=estudio,
            archivo=SimpleUploadedFile("alpha.txt", b"alpha"),
            usuario_subida=self.user,
        )
        doc2 = Documento.objects.create(
            estudio=estudio,
            archivo=SimpleUploadedFile("beta.txt", b"beta"),
            usuario_subida=self.user,
        )
        response = self.client.post(
            "/estudios/documentos/descargar_seleccionados",
            {"doc_id": [str(doc1.id), str(doc2.id)], "id_estudio": str(estudio.id)},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/zip")

    def test_descargar_documento_respetando_id_estudio(self):
        estudio_a = Estudio.objects.create(nombre_estudio="EST-A")
        estudio_b = Estudio.objects.create(nombre_estudio="EST-B")
        doc = Documento.objects.create(
            estudio=estudio_a,
            archivo=SimpleUploadedFile("seguro.txt", b"contenido"),
            usuario_subida=self.user,
        )

        response = self.client.get(f"/estudios/{estudio_b.id}/{doc.id}")

        self.assertEqual(response.status_code, 404)

    def test_descargar_documento_inexistente_en_disco_vuelve_al_repositorio(self):
        estudio = Estudio.objects.create(nombre_estudio="EST-MISS")
        doc = Documento.objects.create(
            estudio=estudio,
            archivo=SimpleUploadedFile("seguro.txt", b"contenido"),
            usuario_subida=self.user,
        )
        os.remove(doc.archivo.path)

        response = self.client.get(f"/estudios/{estudio.id}/{doc.id}", follow=True)

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "no se encuentra disponible")

    def test_repositorio_estudio_muestra_botones_desactivados_sin_seleccion(self):
        estudio = Estudio.objects.create(nombre_estudio="EST-BTN")
        Documento.objects.create(
            estudio=estudio,
            archivo=SimpleUploadedFile("boton.txt", b"contenido"),
            usuario_subida=self.user,
        )

        response = self.client.get(f"/estudios/{estudio.id}")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id="download-button" disabled')
        self.assertContains(response, 'id="delete-button" disabled')

    def test_investigador_no_puede_abrir_estudio_no_asignado_por_url(self):
        investigador = User.objects.create_user(username="investigador", password="secret")
        grupo_investigadores, _ = Group.objects.get_or_create(name="Investigadores")
        investigador.groups.add(grupo_investigadores)
        permisos = Permission.objects.filter(
            codename__in=["can_view_estudios_web", "can_change_estudios_web"]
        )
        investigador.user_permissions.add(*permisos)

        estudio_asignado = Estudio.objects.create(nombre_estudio="EST-ASIG")
        estudio_asignado.investigadores_asociados.add(investigador)
        estudio_ajeno = Estudio.objects.create(nombre_estudio="EST-AJENO")

        self.client.force_login(investigador)

        response = self.client.get(f"/estudios/{estudio_ajeno.id}")

        self.assertEqual(response.status_code, 404)

    def test_investigador_no_puede_descargar_documentos_de_estudio_no_asignado(self):
        investigador = User.objects.create_user(username="investigador_doc", password="secret")
        grupo_investigadores, _ = Group.objects.get_or_create(name="Investigadores")
        investigador.groups.add(grupo_investigadores)
        permisos = Permission.objects.filter(
            codename__in=["can_view_estudios_web", "can_change_estudios_web"]
        )
        investigador.user_permissions.add(*permisos)

        estudio_asignado = Estudio.objects.create(nombre_estudio="EST-DOC-ASIG")
        estudio_asignado.investigadores_asociados.add(investigador)
        estudio_ajeno = Estudio.objects.create(nombre_estudio="EST-DOC-AJENO")
        doc_ajeno = Documento.objects.create(
            estudio=estudio_ajeno,
            archivo=SimpleUploadedFile("ajeno.txt", b"secreto"),
            usuario_subida=self.user,
        )

        self.client.force_login(investigador)

        response = self.client.get(f"/estudios/{estudio_ajeno.id}/{doc_ajeno.id}")

        self.assertEqual(response.status_code, 404)

    def test_investigador_solo_ve_opciones_de_estudios_asignados(self):
        investigador = User.objects.create_user(username="investigador2", password="secret")
        grupo_investigadores, _ = Group.objects.get_or_create(name="Investigadores")
        investigador.groups.add(grupo_investigadores)
        permisos = Permission.objects.filter(codename="can_view_estudios_web")
        investigador.user_permissions.add(*permisos)

        estudio_visible = Estudio.objects.create(
            nombre_estudio="EST-VISIBLE",
            referencia_estudio="REF-VISIBLE",
            investigador_principal="Dra Visible",
        )
        estudio_visible.investigadores_asociados.add(investigador)
        Estudio.objects.create(
            nombre_estudio="EST-OCULTO",
            referencia_estudio="REF-OCULTO",
            investigador_principal="Dr Oculto",
        )

        self.client.force_login(investigador)

        response = self.client.get("/estudios/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "EST-VISIBLE")
        self.assertNotContains(response, "EST-OCULTO")
        self.assertContains(response, "REF-VISIBLE")
        self.assertNotContains(response, "REF-OCULTO")
        self.assertContains(response, "Dra Visible")
        self.assertNotContains(response, "Dr Oculto")

    def test_validacion_excel_limpia_columnas_extra_de_sesiones_previas(self):
        session = self.client.session
        session["columnas_adicionales"] = "Columna antigua"
        session.save()

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append([
            "Referencia del estudio",
            "Nombre del estudio",
            "Descripción",
            "Fecha de inicio",
            "Fecha de fin",
            "Investigador principal",
        ])
        sheet.append(["REF-NEW", "EST-NEW", "Sin extras", "10-01-2026", "12-01-2026", "Dra Test"])
        excel_buffer = io.BytesIO()
        workbook.save(excel_buffer)
        excel_buffer.seek(0)

        response = self.client.post(
            "/estudios/excel",
            {"excel_file": SimpleUploadedFile(
                "estudios.xlsx",
                excel_buffer.getvalue(),
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )},
        )

        self.assertEqual(response.status_code, 200)
        data = self._parse_streaming_json_response(response)
        self.assertEqual(data["status"], "done")

        session = self.client.session
        self.assertNotIn("columnas_adicionales", session)
        self.assertIn("confirmacion_pendiente", session)
        self.assertFalse(session["confirmacion_pendiente"]["context"]["tiene_columnas_extras"])

    def test_descargar_documentos_seleccionados_usa_id_estudio_del_formulario(self):
        estudio_a = Estudio.objects.create(nombre_estudio="EST-POST")
        estudio_b = Estudio.objects.create(nombre_estudio="EST-SESSION")
        doc_a = Documento.objects.create(
            estudio=estudio_a,
            archivo=SimpleUploadedFile("uno.txt", b"uno"),
            usuario_subida=self.user,
        )
        Documento.objects.create(
            estudio=estudio_b,
            archivo=SimpleUploadedFile("dos.txt", b"dos"),
            usuario_subida=self.user,
        )
        session = self.client.session
        session["id"] = estudio_b.id
        session.save()

        response = self.client.post(
            "/estudios/documentos/descargar_seleccionados",
            {"doc_id": [str(doc_a.id)], "id_estudio": str(estudio_a.id)},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/zip")

    def test_eliminar_documento_sin_ids_vuelve_al_repositorio(self):
        estudio = Estudio.objects.create(nombre_estudio="EST-EMPTY")
        session = self.client.session
        session["id"] = estudio.id
        session.save()

        response = self.client.post(
            "/estudios/eliminar_documento",
            {"id_estudio": str(estudio.id)},
        )

        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.url.endswith(f"/estudios/{estudio.id}"))

    def test_crear_recurso_lo_guarda_como_definitivo(self):
        response = self.client.post(
            "/archivo/crear_recurso",
            {
                "congelador": "FZ-NUEVO",
                "modelo": "Model X",
                "ano_compra": 2024,
                "temperatura_objetivo": "-80",
                "localizacion_edificio": "Edificio A",
                "tipo_estructura": "congelador_80",
            },
        )

        self.assertEqual(response.status_code, 302)
        recurso = Congelador.objects.get(congelador="FZ-NUEVO")
        self.assertEqual(recurso.tipo_estructura, "congelador_80")
        self.assertTrue(response.url.endswith(f"/archivo/detalles_congelador/{recurso.congelador}"))

    def test_crear_recurso_requiere_tipo_estructura(self):
        response = self.client.post(
            "/archivo/crear_recurso",
            {
                "congelador": "FZ-SIN-TIPO",
                "modelo": "Model X",
                "ano_compra": 2024,
                "temperatura_objetivo": "-80",
                "localizacion_edificio": "Edificio A",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Este campo es obligatorio")
        self.assertFalse(Congelador.objects.filter(congelador="FZ-SIN-TIPO").exists())

    def test_formulario_editar_recurso_bloquea_tipo_estructura(self):
        recurso = Congelador.objects.create(
            congelador="FZ-BLOQ",
            tipo_estructura="congelador_80",
        )

        form = CongeladorForm(instance=recurso)

        self.assertTrue(form.fields["tipo_estructura"].disabled)

    def test_usuario_con_solo_add_localizaciones_no_puede_editar_recurso(self):
        editor_limitado = User.objects.create_user(username="soloadd", password="secret")
        permiso_add = Permission.objects.get(codename="can_add_localizaciones_web")
        permiso_view = Permission.objects.get(codename="can_view_localizaciones_web")
        editor_limitado.user_permissions.add(permiso_add, permiso_view)
        self.client.force_login(editor_limitado)

        response = self.client.get(
            f"/archivo/detalles_congelador/{self.congelador.congelador}/editar"
        )

        self.assertEqual(response.status_code, 302)
        self.assertIn("/accounts/login/", response.url)

    def test_detalle_recurso_oculta_acciones_de_cambio_sin_permiso(self):
        visor = User.objects.create_user(username="visorrec", password="secret")
        permiso_view = Permission.objects.get(codename="can_view_localizaciones_web")
        visor.user_permissions.add(permiso_view)
        reparacion = ReparacionRecurso.objects.create(
            recurso=self.congelador,
            fecha_reparacion=date(2026, 1, 15),
            descripcion="Revision general",
        )
        self.client.force_login(visor)

        response = self.client.get(f"/archivo/detalles_congelador/{self.congelador.congelador}")

        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "Editar detalles del recurso")
        self.assertNotContains(response, "Anadir reparacion")
        self.assertNotContains(response, f'/archivo/reparaciones/{reparacion.id}/editar')
        self.assertNotContains(response, f'/archivo/reparaciones/{reparacion.id}/eliminar')

    def test_localizaciones_oculta_acciones_restringidas_sin_permiso(self):
        visor = User.objects.create_user(username="visorarchivo", password="secret")
        permiso_view = Permission.objects.get(codename="can_view_localizaciones_web")
        visor.user_permissions.add(permiso_view)
        self.client.force_login(visor)

        response = self.client.get("/archivo/")

        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "Crear recurso")
        self.assertNotContains(response, 'id="eliminar-btn"')

    def test_detalles_congelador_requiere_autenticacion(self):
        self.client.logout()

        response = self.client.get(f"/archivo/detalles_congelador/{self.congelador.congelador}")

        self.assertEqual(response.status_code, 302)
        self.assertIn("/login", response.url)

    def test_rutas_principales_redirigen_a_login_si_no_hay_sesion(self):
        muestra = Muestra.objects.create(nom_lab="LAB-AUTH", estado_actual="DISP")
        self.subposicion_libre.muestra = muestra
        self.subposicion_libre.vacia = False
        self.subposicion_libre.save()

        rutas = [
            "/",
            "/muestras/nueva",
            f"/archivo/detalles_muestra/{muestra.nom_lab}",
            f"/muestras/historial_localizaciones/{muestra.id}",
            "/archivo/crear_recurso",
            f"/archivo/plantilla_recurso/{self.congelador.id}",
            "/archivo/nuevo/0",
            "/estudios/excel",
        ]

        self.client.logout()

        for ruta in rutas:
            with self.subTest(ruta=ruta):
                response = self.client.get(ruta)
                self.assertEqual(response.status_code, 302)
                self.assertIn("/accounts/login/", response.url)

    def test_login_redirige_a_inicio_si_el_usuario_ya_esta_autenticado(self):
        response = self.client.get("/accounts/login/")

        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, "/")

    def test_logout_redirige_a_login_con_confirmacion(self):
        response = self.client.post("/accounts/logout/")

        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, "/accounts/login/?logged_out=1")

        follow_response = self.client.get(response.url, follow=True)
        self.assertEqual(follow_response.status_code, 200)
        self.assertContains(follow_response, "La sesion se ha cerrado correctamente.")

    def test_login_con_next_redirige_a_la_ruta_original(self):
        self.client.logout()
        destino = "/archivo/"

        response = self.client.post(
            f"/accounts/login/?next={quote(destino)}",
            {"username": "tester", "password": "secret", "next": destino},
        )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, destino)

    def test_form_reparacion_rechaza_fecha_anterior_al_ano_compra(self):
        recurso = Congelador.objects.create(congelador="FZ-REPARA", ano_compra=2024)

        form = ReparacionRecursoForm(
            data={
                "fecha_reparacion": "2023-12-31",
                "descripcion": "Cambio de compresor",
                "coste": "120.50",
            },
            recurso=recurso,
        )

        self.assertFalse(form.is_valid())
        self.assertIn("fecha_reparacion", form.errors)

    def test_model_reparacion_rechaza_fecha_anterior_al_ano_compra(self):
        recurso = Congelador.objects.create(congelador="FZ-MODEL", ano_compra=2024)
        reparacion = ReparacionRecurso(
            recurso=recurso,
            fecha_reparacion=date(2023, 1, 1),
            descripcion="Revision inicial",
        )

        with self.assertRaises(ValidationError):
            reparacion.full_clean()

    def test_exportar_congeladores_devuelve_zip_por_tipos(self):
        response = self.client.get("/archivo/exportar_congeladores")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/zip")
        with zipfile.ZipFile(io.BytesIO(response.content)) as zip_file:
            nombres = sorted(zip_file.namelist())
            self.assertEqual(
                nombres,
                [
                    "recursos_congelador_-20.xlsx",
                    "recursos_congelador_-80.xlsx",
                    "recursos_neveras.xlsx",
                ],
            )

            workbook_20 = openpyxl.load_workbook(io.BytesIO(zip_file.read("recursos_congelador_-20.xlsx")))
            sheet_20 = workbook_20.active
            self.assertEqual(sheet_20.cell(row=2, column=1).value, "FZ-20")
            self.assertEqual(sheet_20.cell(row=2, column=4).value, "C1")
            self.assertEqual(sheet_20.cell(row=2, column=5).value, "BOX-20")
            self.assertEqual(sheet_20.cell(row=2, column=6).value, "A")

            workbook_80 = openpyxl.load_workbook(io.BytesIO(zip_file.read("recursos_congelador_-80.xlsx")))
            sheet_80 = workbook_80.active
            self.assertEqual(sheet_80.cell(row=2, column=1).value, "FZ-01")
            self.assertEqual(sheet_80.cell(row=2, column=8).value, "C1")

    def test_exportar_congeladores_seleccionados_devuelve_zip_por_recurso(self):
        response = self.client.get(
            f"/archivo/exportar_congeladores_seleccionados?congelador={self.congelador.id}&congelador={self.congelador_20.id}"
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/zip")
        with zipfile.ZipFile(io.BytesIO(response.content)) as zip_file:
            nombres = sorted(zip_file.namelist())
            self.assertEqual(nombres, ["recurso_fz-01.xlsx", "recurso_fz-20.xlsx"])

    def test_exportar_posiciones_libres_devuelve_zip_por_recurso(self):
        response = self.client.get(
            f"/archivo/exportar_posiciones_libres?congelador={self.congelador.id}&congelador={self.nevera.id}"
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/zip")
        with zipfile.ZipFile(io.BytesIO(response.content)) as zip_file:
            nombres = sorted(zip_file.namelist())
            self.assertEqual(nombres, ["posiciones_libres_fz-01.xlsx", "posiciones_libres_nv-01.xlsx"])

            workbook_80 = openpyxl.load_workbook(io.BytesIO(zip_file.read("posiciones_libres_fz-01.xlsx")))
            sheet_80 = workbook_80.active
            self.assertEqual(sheet_80.cell(row=2, column=1).value, "FZ-01")
            self.assertEqual(sheet_80.cell(row=2, column=9).value, "A")

            workbook_nevera = openpyxl.load_workbook(io.BytesIO(zip_file.read("posiciones_libres_nv-01.xlsx")))
            sheet_nevera = workbook_nevera.active
            self.assertEqual(sheet_nevera.cell(row=2, column=1).value, "NV-01")
            self.assertEqual(sheet_nevera.cell(row=2, column=5).value, "BOX-NV")
            self.assertEqual(sheet_nevera.cell(row=2, column=6).value, "B")

    def test_eliminar_estante_con_muestras_devuelve_mensaje_claro(self):
        muestra = Muestra.objects.create(
            nom_lab="LAB-BLOQ",
            estado_actual="DISP",
        )
        self.subposicion_libre.muestra = muestra
        self.subposicion_libre.vacia = False
        self.subposicion_libre.save()

        response = self.client.post(
            "/archivo/eliminar_localizacion",
            {"estante": [str(self.estante.id)]},
        )

        self.assertEqual(response.status_code, 200)
        data = self._parse_streaming_json_response(response)
        self.assertEqual(data["status"], "error")
        self.assertIn('No se puede eliminar el estante "1"', data["message"])
        self.assertIn('LAB-BLOQ', data["message"])
        self.assertIn('Caja C1', data["message"])

    def test_eliminar_caja_con_muestras_devuelve_mensaje_claro(self):
        muestra = Muestra.objects.create(
            nom_lab="LAB-BOX",
            estado_actual="DISP",
        )
        self.subposicion_libre.muestra = muestra
        self.subposicion_libre.vacia = False
        self.subposicion_libre.save()

        response = self.client.post(
            "/archivo/eliminar_localizacion",
            {"caja": [str(self.caja.id)]},
        )

        self.assertEqual(response.status_code, 200)
        data = self._parse_streaming_json_response(response)
        self.assertEqual(data["status"], "error")
        self.assertIn('No se puede eliminar la caja "C1"', data["message"])
        self.assertIn('Subposición A-1', data["message"])

    def test_opciones_creacion_localizaciones_muestra_los_dos_caminos(self):
        response = self.client.get(f"/archivo/detalles_congelador/{self.congelador.congelador}/crear_posiciones")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Mediante formulario")
        self.assertContains(response, "Mediante plantilla Excel")

    def test_confirmacion_upload_localizaciones_muestra_preview_si_no_hay_errores(self):
        session = self.client.session
        session["confirmacion_pendiente"] = {
            "template": "confirmacion_upload_localizacion.html",
            "context": {
                "errores_encontrados": False,
                "preview_recurso": {
                    "congelador": self.congelador_20.congelador,
                    "simple_structure": True,
                    "summary_total_posiciones": 25,
                    "summary_libres": 25,
                    "summary_ocupadas": 0,
                    "summary_cajas": 1,
                    "estantes": [
                        {
                            "numero": "E1",
                            "racks": [
                                {
                                    "numero": "C1",
                                    "posicion_rack_estante": "1",
                                    "numero_cajas": 1,
                                }
                            ],
                        }
                    ],
                },
            },
            "mensajes": [{"level": "success", "text": "No tiene errores en ningun campo."}],
        }
        session.save()

        response = self.client.get(f"/archivo/nuevo?mostrar_confirmacion=1&congelador_id={self.congelador_20.id}")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Vista previa de la estructura que se va a crear")
        self.assertContains(response, "Recurso FZ-20")
        self.assertContains(response, "Cajon C1 - Posicion dentro del estante: 1")
        self.assertNotContains(response, "25 libres")
        self.assertNotContains(response, "0 ocupadas")
        self.assertNotContains(response, "Importar otro archivo Excel")

    def test_formulario_dinamico_muestra_inputs_y_formatos_de_caja(self):
        response = self.client.get(
            f"/archivo/detalles_congelador/{self.congelador.congelador}/crear_posiciones/formulario"
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'placeholder="Nombre del estante"')
        self.assertContains(response, "5x5 (25 posiciones)")
        self.assertContains(response, "9x9 (81 posiciones)")
        self.assertContains(response, "10x10 (100 posiciones)")

    def test_importacion_muestras_excel_muestra_paso_inicial_de_seleccion(self):
        response = self.client.get("/muestras/upload_excel")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Paso 1 de 2")
        self.assertContains(response, "Solo datos de la muestra")
        self.assertContains(response, "ubicación en congelador -80")

    def test_importacion_muestras_excel_muestra_instrucciones_segun_tipo(self):
        response = self.client.get("/muestras/upload_excel?import_mode=congelador_80")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Paso 2 de 2")
        self.assertContains(response, "estructura completa de congelador -80")
        self.assertContains(response, "Congelador, Estante, Rack, Bandeja, Caja, Fila y Columna")
        self.assertContains(response, 'value="congelador_80"')

    def test_crear_localizacion_manual_simple_crea_subposicion(self):
        structure = {
            "estantes": [
                {
                    "name": "E2",
                    "racks": [
                        {
                            "name": "CJ-NEW",
                            "cajas": [
                                {"name": "BOX-NEW", "layout": "5x5"}
                            ],
                        }
                    ],
                }
            ]
        }
        response = self.client.post(
            f"/archivo/detalles_congelador/{self.congelador_20.congelador}/crear_posiciones/formulario",
            {
                "structure_json": json.dumps(structure),
            },
        )

        self.assertEqual(response.status_code, 302)
        self.assertTrue(
            Subposicion.objects.filter(
                caja__rack__estante__congelador=self.congelador_20,
                caja__numero="BOX-NEW",
                fila="C",
                columna="4",
            ).exists()
        )

    def test_crear_localizacion_manual_80_genera_subposiciones_segun_formato(self):
        structure = {
            "estantes": [
                {
                    "name": "E3",
                    "racks": [
                        {
                            "name": "R-01",
                            "bandejas": [
                                {
                                    "name": "B-01",
                                    "cajas": [
                                        {"name": "C-01", "layout": "10x10"}
                                    ],
                                }
                            ],
                        }
                    ],
                }
            ]
        }

        response = self.client.post(
            f"/archivo/detalles_congelador/{self.congelador.congelador}/crear_posiciones/formulario",
            {"structure_json": json.dumps(structure)},
        )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(
            Subposicion.objects.filter(caja__numero="C-01", caja__rack__estante__congelador=self.congelador).count(),
            100,
        )
        self.assertTrue(
            Subposicion.objects.filter(
                caja__numero="C-01",
                caja__rack__estante__congelador=self.congelador,
                fila="J",
                columna="10",
            ).exists()
        )

    def test_formulario_manual_precarga_estructura_existente_como_solo_lectura(self):
        response = self.client.get(
            f"/archivo/detalles_congelador/{self.congelador.congelador}/crear_posiciones/formulario"
        )

        self.assertEqual(response.status_code, 200)
        initial_data = response.context["initial_structure_json"]
        self.assertEqual(initial_data["estantes"][0]["name"], "1")
        self.assertTrue(initial_data["estantes"][0]["existing"])
        self.assertTrue(initial_data["estantes"][0]["racks"][0]["existing"])
        self.assertTrue(initial_data["estantes"][0]["racks"][0]["bandejas"][0]["cajas"][0]["existing"])
        self.assertContains(response, "Ya creado")
        self.assertContains(response, "Ya creada")

    def test_crear_localizacion_manual_permita_ampliar_estructura_existente(self):
        structure = {
            "estantes": [
                {
                    "name": "1",
                    "existing": True,
                    "racks": [
                        {
                            "name": "A",
                            "existing": True,
                            "bandejas": [
                                {
                                    "name": "B1",
                                    "existing": True,
                                    "cajas": [
                                        {"name": "C1", "layout": "9x9", "existing": True},
                                        {"name": "C2", "layout": "5x5"},
                                        {"name": "C3", "layout": "5x5"},
                                    ],
                                }
                            ],
                        }
                    ],
                }
            ]
        }

        response = self.client.post(
            f"/archivo/detalles_congelador/{self.congelador.congelador}/crear_posiciones/formulario",
            {"structure_json": json.dumps(structure)},
        )

        self.assertEqual(response.status_code, 302)
        self.assertTrue(
            Caja.objects.filter(
                rack=self.rack,
                bandeja=self.bandeja,
                numero="C2",
                posicion_caja_rack="2",
            ).exists()
        )
        self.assertTrue(
            Caja.objects.filter(
                rack=self.rack,
                bandeja=self.bandeja,
                numero="C3",
                posicion_caja_rack="3",
            ).exists()
        )
        self.assertEqual(
            Subposicion.objects.filter(caja__rack__estante__congelador=self.congelador, caja__numero="C1").count(),
            1,
        )
        self.assertEqual(
            Subposicion.objects.filter(caja__rack__estante__congelador=self.congelador, caja__numero="C2").count(),
            25,
        )
        self.assertEqual(
            Subposicion.objects.filter(caja__rack__estante__congelador=self.congelador, caja__numero="C3").count(),
            25,
        )

    def test_crear_localizacion_manual_rechaza_estantes_duplicados(self):
        structure = {
            "estantes": [
                {"name": "E2", "racks": []},
                {"name": "E2", "racks": []},
            ]
        }

        response = self.client.post(
            f"/archivo/detalles_congelador/{self.congelador_20.congelador}/crear_posiciones/formulario",
            {"structure_json": json.dumps(structure)},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(Estante.objects.filter(congelador=self.congelador_20, numero="E2").count(), 0)
        self.assertContains(response, "No puede haber dos estantes con el mismo nombre.")

    def test_anadir_muestras_muestra_selector_de_recurso_compatible_con_todas_las_estructuras(self):
        response = self.client.get("/muestras/nueva")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Recurso")
        self.assertContains(response, 'data-structure="congelador_80"')
        self.assertContains(response, 'data-structure="congelador_20"')
        self.assertContains(response, 'data-structure="nevera"')
        self.assertContains(response, 'id="subposicion-grid-shell"')

@override_settings(MEDIA_ROOT=TEST_MEDIA_ROOT)
class GesLabFileCleanupTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="filetester", password="secret")

    def test_eliminar_documento_borra_el_archivo_fisico(self):
        estudio = Estudio.objects.create(nombre_estudio="EST-FILE")
        documento = Documento.objects.create(
            estudio=estudio,
            archivo=SimpleUploadedFile("doc.txt", b"contenido"),
            usuario_subida=self.user,
        )

        ruta_archivo = documento.archivo.path
        self.assertTrue(os.path.exists(ruta_archivo))

        documento.delete()

        self.assertFalse(os.path.exists(ruta_archivo))

    def test_reemplazar_o_borrar_fotografia_limpia_el_archivo_antiguo(self):
        recurso = Congelador.objects.create(
            congelador="FZ-FOTO",
            tipo_estructura=Congelador.ESTRUCTURA_CONGELADOR_80,
            fotografia=SimpleUploadedFile(
                "foto1.jpg",
                b"\xff\xd8\xff\xe0" + b"0" * 32,
                content_type="image/jpeg",
            ),
        )

        ruta_inicial = recurso.fotografia.path
        self.assertTrue(os.path.exists(ruta_inicial))

        recurso.fotografia = SimpleUploadedFile(
            "foto2.jpg",
            b"\xff\xd8\xff\xe0" + b"1" * 32,
            content_type="image/jpeg",
        )
        recurso.save()

        self.assertFalse(os.path.exists(ruta_inicial))
        ruta_nueva = recurso.fotografia.path
        self.assertTrue(os.path.exists(ruta_nueva))

        recurso.delete()

        self.assertFalse(os.path.exists(ruta_nueva))
