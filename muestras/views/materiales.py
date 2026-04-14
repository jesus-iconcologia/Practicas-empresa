import base64
import io
import os

import openpyxl
import pandas as pd
from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required, permission_required
from django.core.exceptions import PermissionDenied
from django.core.paginator import EmptyPage, PageNotAnInteger, Paginator
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from openpyxl.styles import PatternFill

from ..forms import MaterialForm, UploadExcel
from ..models import Material
from ..parameters_config import get_excel_colors


def _excel_text_or_none(value):
    if value is None or pd.isna(value):
        return None
    text = str(value).strip()
    return text or None


def _parse_material_bool(value):
    text = _excel_text_or_none(value)
    if text is None:
        return None
    lowered = text.lower()
    if lowered in {"1", "true", "si", "sí", "x", "yes"}:
        return True
    if lowered in {"0", "false", "no", "n", "off"}:
        return False
    return "invalid"


@login_required
@permission_required("muestras.can_view_muestras_web")
def materiales_todos(request):
    materiales = Material.objects.all().order_by("nombre")

    nombre = request.GET.get("nombre", "").strip()
    if nombre:
        materiales = materiales.filter(nombre__icontains=nombre)

    usa_volumen = request.GET.get("usa_volumen", "").strip()
    if usa_volumen == "si":
        materiales = materiales.filter(usa_volumen=True)
    elif usa_volumen == "no":
        materiales = materiales.filter(usa_volumen=False)

    usa_concentracion = request.GET.get("usa_concentracion", "").strip()
    if usa_concentracion == "si":
        materiales = materiales.filter(usa_concentracion=True)
    elif usa_concentracion == "no":
        materiales = materiales.filter(usa_concentracion=False)

    usa_unidades = request.GET.get("usa_unidades", "").strip()
    if usa_unidades == "si":
        materiales = materiales.filter(usa_unidades_bloque=True)
    elif usa_unidades == "no":
        materiales = materiales.filter(usa_unidades_bloque=False)

    busqueda = request.GET.get("busqueda", "").strip()
    if busqueda:
        materiales = materiales.filter(nombre__icontains=busqueda)

    contador_total = materiales.count()
    items_por_pagina = request.GET.get("items_por_pagina", 25)
    if str(items_por_pagina) == "todas":
        items_por_pagina = "todas"
        paginator = Paginator(materiales, max(contador_total, 1))
    else:
        try:
            items_por_pagina = int(items_por_pagina)
            if items_por_pagina not in [10, 25, 50, 100]:
                items_por_pagina = 25
        except (ValueError, TypeError):
            items_por_pagina = 25
        paginator = Paginator(materiales, items_por_pagina)

    numero_pagina = request.GET.get("page", 1)
    try:
        materiales_page = paginator.page(numero_pagina)
    except PageNotAnInteger:
        materiales_page = paginator.page(1)
    except EmptyPage:
        materiales_page = paginator.page(paginator.num_pages)

    return render(
        request,
        "materiales_todos.html",
        {
            "materiales": materiales_page.object_list,
            "materiales_page": materiales_page,
            "paginator": paginator,
            "contador_materiales": contador_total,
            "items_por_pagina": items_por_pagina,
            "busqueda": busqueda,
            "request": request,
        },
    )


@login_required
@permission_required("muestras.can_view_muestras_web")
def acciones_materiales(request):
    if request.method != "POST":
        return redirect("materiales_todos")

    materiales_ids = request.POST.getlist("material_id")
    queryset = Material.objects.filter(id__in=materiales_ids)

    if "exportar_seleccionados" in request.POST:
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        if not materiales_ids:
            messages.error(request, "No has seleccionado ningún material para exportar.")
            return redirect("materiales_todos")

        response["Content-Disposition"] = 'attachment; filename="materiales_seleccionados.xlsx"'

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Materiales"
        headers = [
            "Nombre del material",
            "Usa volumen",
            "Unidades de volumen",
            "Usa concentración",
            "Unidades de concentración",
            "Usa unidades",
            "Unidades",
        ]
        for col, header in enumerate(headers, 1):
            ws.cell(row=1, column=col).value = header

        row_num = 2
        for material in queryset.order_by("nombre"):
            ws.cell(row=row_num, column=1).value = material.nombre
            ws.cell(row=row_num, column=2).value = "Sí" if material.usa_volumen else "No"
            ws.cell(row=row_num, column=3).value = material.unidades_volumen or ""
            ws.cell(row=row_num, column=4).value = "Sí" if material.usa_concentracion else "No"
            ws.cell(row=row_num, column=5).value = material.unidades_concentracion or ""
            ws.cell(row=row_num, column=6).value = "Sí" if material.usa_unidades_bloque else "No"
            ws.cell(row=row_num, column=7).value = material.unidades_recuento or ""
            row_num += 1

        wb.save(response)
        return response

    if "eliminar" in request.POST:
        if not request.user.has_perm("muestras.can_delete_muestras_web"):
            raise PermissionDenied
        if not materiales_ids:
            messages.error(request, "No has seleccionado ningún material para eliminar.")
            return redirect("materiales_todos")

        materiales_con_uso = list(queryset.filter(muestras__isnull=False).distinct())
        if materiales_con_uso:
            nombres = ", ".join(material.nombre for material in materiales_con_uso[:5])
            messages.error(
                request,
                f"No se pueden eliminar materiales que ya están asociados a muestras: {nombres}.",
            )
            return redirect("materiales_todos")

        total = queryset.count()
        queryset.delete()
        messages.success(request, f"{total} materiales eliminados correctamente.")
        return redirect("materiales_todos")

    return redirect("materiales_todos")


@login_required
@permission_required("muestras.can_add_muestras_web")
def upload_excel_materiales(request):
    if request.GET.get("mostrar_confirmacion") and "confirmacion_materiales_pendiente" in request.session:
        datos_conf = request.session.pop("confirmacion_materiales_pendiente")
        for msg in datos_conf.get("mensajes", []):
            getattr(messages, msg["level"])(request, msg["text"])
        return render(request, datos_conf["template"], datos_conf.get("context", {}))

    if request.method == "POST":
        form = UploadExcel(request.POST, request.FILES)
        if "confirmar" in request.POST:
            filas_validas = request.session.get("filas_validas_materiales", [])
            with transaction.atomic():
                for datos in filas_validas:
                    Material.objects.create(
                        nombre=datos["nombre"],
                        usa_volumen=datos["usa_volumen"],
                        unidades_volumen=datos["unidades_volumen"],
                        usa_concentracion=datos["usa_concentracion"],
                        unidades_concentracion=datos["unidades_concentracion"],
                        usa_unidades_bloque=datos["usa_unidades_bloque"],
                        unidades_recuento=datos["unidades_recuento"],
                        activo=True,
                    )
            request.session.pop("filas_validas_materiales", None)
            request.session.pop("errores_materiales", None)
            request.session.pop("excel_materiales_base64", None)
            request.session.pop("columnas_adicionales_materiales", None)
            messages.success(request, "Los materiales se han creado correctamente.")
            return redirect("materiales_todos")

        if "cancelar" in request.POST:
            request.session.pop("filas_validas_materiales", None)
            request.session.pop("errores_materiales", None)
            request.session.pop("excel_materiales_base64", None)
            request.session.pop("columnas_adicionales_materiales", None)
            messages.error(request, "La importación de materiales se ha cancelado.")
            return redirect("upload_excel_materiales")

        if "excel_errores" in request.POST:
            errores = request.session.get("errores_materiales", {})
            excel_b64 = request.session.get("excel_materiales_base64")
            if not excel_b64:
                messages.error(request, "No hay un Excel validado para descargar.")
                return redirect("upload_excel_materiales")

            wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(excel_b64)))
            ws = wb.active
            colors = get_excel_colors()
            error_fill = PatternFill("solid", fgColor=colors["error_row"])

            for fila, detalle in errores.items():
                if detalle.get("bloqueantes"):
                    for cell in ws[fila]:
                        cell.fill = error_fill

            response = HttpResponse(content_type="application/ms-excel")
            response["Content-Disposition"] = 'attachment; filename="materiales_con_errores.xlsx"'
            wb.save(response)
            return response

        if form.is_valid():
            excel_file = request.FILES["excel_file"]
            excel_bytes = excel_file.read()
            request.session["excel_materiales_base64"] = base64.b64encode(excel_bytes).decode()
            df = pd.read_excel(io.BytesIO(excel_bytes))

            rename_columns = {
                "nombre del material": "nombre",
                "usa volumen": "usa_volumen",
                "unidades de volumen": "unidades_volumen",
                "usa concentración": "usa_concentracion",
                "unidades de concentración": "unidades_concentracion",
                "usa unidades": "usa_unidades_bloque",
                "unidades": "unidades_recuento",
            }
            df.columns = df.columns.str.strip().str.lower()
            columnas_esperadas = set(rename_columns.keys())
            columnas_existentes = set(df.columns)
            columnas_faltantes = columnas_esperadas - columnas_existentes

            if columnas_faltantes:
                columnas_str = ", ".join(sorted(columnas_faltantes))
                return render(
                    request,
                    "upload_excel_materiales.html",
                    {"form": form, "error": f"Error de formato: faltan columnas esperadas: {columnas_str}"},
                )

            if df.empty or len(df) == 0:
                return render(
                    request,
                    "upload_excel_materiales.html",
                    {
                        "form": form,
                        "error": "Error de formato: el archivo Excel está vacío o no contiene filas de datos.",
                    },
                )

            columnas_adicionales = columnas_existentes - columnas_esperadas
            columnas_adicionales_str = ", ".join(sorted(columnas_adicionales))
            request.session["columnas_adicionales_materiales"] = columnas_adicionales_str

            df = df.rename(columns=rename_columns)
            errores = {}
            filas_validas = []
            nombres_excel = set()
            nombres_existentes = {nombre.lower() for nombre in Material.objects.values_list("nombre", flat=True)}

            for idx, row in df.iterrows():
                fila = idx + 2
                errores[fila] = {"bloqueantes": []}

                nombre = _excel_text_or_none(row.get("nombre"))
                usa_volumen = _parse_material_bool(row.get("usa_volumen"))
                usa_concentracion = _parse_material_bool(row.get("usa_concentracion"))
                usa_unidades_bloque = _parse_material_bool(row.get("usa_unidades_bloque"))

                datos = {
                    "nombre": nombre,
                    "usa_volumen": bool(usa_volumen is True),
                    "unidades_volumen": _excel_text_or_none(row.get("unidades_volumen")),
                    "usa_concentracion": bool(usa_concentracion is True),
                    "unidades_concentracion": _excel_text_or_none(row.get("unidades_concentracion")),
                    "usa_unidades_bloque": bool(usa_unidades_bloque is True),
                    "unidades_recuento": _excel_text_or_none(row.get("unidades_recuento")),
                }

                if not nombre:
                    errores[fila]["bloqueantes"].append("campo_obligatorio_vacio:nombre")
                else:
                    nombre_lower = nombre.lower()
                    if nombre_lower in nombres_existentes:
                        errores[fila]["bloqueantes"].append("material_existente_bd")
                    if nombre_lower in nombres_excel:
                        errores[fila]["bloqueantes"].append("material_duplicado_excel")
                    else:
                        nombres_excel.add(nombre_lower)

                if usa_volumen == "invalid":
                    errores[fila]["bloqueantes"].append("valor_invalido:usa_volumen")
                if usa_concentracion == "invalid":
                    errores[fila]["bloqueantes"].append("valor_invalido:usa_concentracion")
                if usa_unidades_bloque == "invalid":
                    errores[fila]["bloqueantes"].append("valor_invalido:usa_unidades")

                if not errores[fila]["bloqueantes"]:
                    filas_validas.append(datos)

            request.session["filas_validas_materiales"] = filas_validas
            request.session["errores_materiales"] = errores

            numero_errores_bloqueantes = sum(1 for detalle in errores.values() if detalle["bloqueantes"])
            errores_encontrados = numero_errores_bloqueantes > 0 or bool(columnas_adicionales)
            mensajes = [{"level": "info", "text": f"El excel contiene {len(df)} registros."}]
            if numero_errores_bloqueantes == 0 and not columnas_adicionales:
                mensajes.append({"level": "success", "text": "No tiene errores en ningún campo."})
            else:
                if numero_errores_bloqueantes > 0:
                    mensajes.append(
                        {
                            "level": "error",
                            "text": f"Contiene {numero_errores_bloqueantes} filas con errores graves.",
                        }
                    )
                if columnas_adicionales:
                    mensajes.append(
                        {
                            "level": "warning",
                            "text": f"Contiene {len(columnas_adicionales)} columnas extras: {columnas_adicionales_str}",
                        }
                    )

            request.session["confirmacion_materiales_pendiente"] = {
                "template": "confirmacion_upload_materiales.html",
                "context": {
                    "errores_encontrados": errores_encontrados,
                    "numero_errores_bloqueantes": numero_errores_bloqueantes,
                    "tiene_columnas_extras": bool(columnas_adicionales),
                },
                "mensajes": mensajes,
            }
            request.session.save()
            return redirect(f"{request.path}?mostrar_confirmacion=1")
    else:
        form = UploadExcel()

    return render(request, "upload_excel_materiales.html", {"form": form})


@login_required
@permission_required("muestras.can_view_muestras_web")
def descargar_plantilla_materiales(request):
    plantilla_path = os.path.join(
        settings.BASE_DIR,
        "geslab_project",
        "globalstaticfiles",
        "files",
        "plantilla_materiales.xlsx",
    )
    if not os.path.exists(plantilla_path):
        return HttpResponse("La plantilla no se encuentra disponible.", status=404)

    with open(plantilla_path, "rb") as plantilla_file:
        response = HttpResponse(
            plantilla_file.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="plantilla_materiales.xlsx"'
        return response


@login_required
@permission_required("muestras.can_add_muestras_web")
def nuevo_material(request):
    if request.method == "POST":
        form = MaterialForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Se ha creado el material con éxito.")
            return redirect("materiales_todos")
    else:
        form = MaterialForm()
    return render(request, "nuevo_material.html", {"form": form})


@login_required
@permission_required("muestras.can_change_muestras_web")
def editar_material(request, material_id):
    material = get_object_or_404(Material, pk=material_id)
    if request.method == "POST":
        form = MaterialForm(request.POST, instance=material)
        if form.is_valid():
            form.save()
            messages.success(request, "Material actualizado correctamente.")
            return redirect("materiales_todos")
    else:
        form = MaterialForm(instance=material)
    return render(request, "editar_material.html", {"form": form, "material": material})


@login_required
@permission_required("muestras.can_delete_muestras_web")
def eliminar_material(request, material_id):
    if request.method != "POST":
        return redirect("materiales_todos")
    material = get_object_or_404(Material, pk=material_id)
    if material.muestras.exists():
        messages.error(request, "No se puede eliminar un material que ya está asociado a muestras.")
        return redirect("materiales_todos")
    material.delete()
    messages.success(request, "Material eliminado correctamente.")
    return redirect("materiales_todos")
